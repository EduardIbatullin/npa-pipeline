"""Тесты точки входа OCR → DOCX (гибридная схема, см. docs/stage3-plan.md и
модуль npa_pipeline.ocr). Движки — фейковые функции на реальных фикстурах/данных,
без загрузки тяжёлых моделей."""

from __future__ import annotations

import json
from pathlib import Path

import pytest

from npa_pipeline.ocr import (
    UNRECOGNIZABLE_MARK,
    OcrDocumentResult,
    PageResult,
    TableBlock,
    TextBlock,
    build_docx,
    expected_signatory_title,
    ocr_document,
    ocr_page,
    run_paddle_tables,
    run_tesseract_text,
)
from npa_pipeline.ocr_table import TableCell

FIXTURES = Path(__file__).parent / "fixtures" / "ocr_tables"


def _load_simple_table_json() -> dict:
    """Фикстура в fixtures/ocr_tables/ хранится в "плоском" виде (для прямых тестов
    reconstruct_table в test_ocr_table.py) — rec_boxes/rec_texts не вложены в
    table_ocr_pred, как в реальном JSON движка. Здесь оборачиваем обратно в форму,
    которую реально отдаёт PPStructureV3 (см. _extract_table_blocks)."""
    with open(FIXTURES / "simple_table_0.json", encoding="utf-8") as f:
        flat = json.load(f)
    nested = {
        "cell_box_list": flat["cell_box_list"],
        "table_ocr_pred": {
            "rec_boxes": flat["rec_boxes"],
            "rec_texts": flat["rec_texts"],
        },
    }
    return {"table_res_list": [nested]}


def _tesseract_data(words: list[tuple[str, int, int, int, int, int]]) -> dict:
    """words — (text, conf, left, top, width, height); остальные поля Tesseract
    (block/par/line_num) проставляются по порядку — одна строка на все слова,
    если явно не указано иное, этого достаточно для тестов группировки."""
    data = {
        "text": [],
        "conf": [],
        "left": [],
        "top": [],
        "width": [],
        "height": [],
        "block_num": [],
        "par_num": [],
        "line_num": [],
    }
    for text, conf, left, top, width, height in words:
        data["text"].append(text)
        data["conf"].append(conf)
        data["left"].append(left)
        data["top"].append(top)
        data["width"].append(width)
        data["height"].append(height)
        data["block_num"].append(1)
        data["par_num"].append(1)
        data["line_num"].append(1)
    return data


def _tesseract_data_multiline(words: list[tuple[str, int, int, int, int, int, int]]) -> dict:
    """Как _tesseract_data, но с явным line_num (7-й элемент каждого тьюпла) —
    нужно воспроизвести реальный случай, когда Tesseract разбивает одну визуальную
    строку на несколько геометрических из-за помехи печати."""
    data = {
        "text": [],
        "conf": [],
        "left": [],
        "top": [],
        "width": [],
        "height": [],
        "block_num": [],
        "par_num": [],
        "line_num": [],
    }
    for text, conf, left, top, width, height, line_num in words:
        data["text"].append(text)
        data["conf"].append(conf)
        data["left"].append(left)
        data["top"].append(top)
        data["width"].append(width)
        data["height"].append(height)
        data["block_num"].append(1)
        data["par_num"].append(1)
        data["line_num"].append(line_num)
    return data


def test_run_paddle_tables_real_fixture():
    fixture = _load_simple_table_json()
    tables = run_paddle_tables(Path("unused.png"), lambda _p: fixture)
    assert len(tables) == 1
    t = tables[0]
    assert t.n_rows == 7
    assert t.n_cols == 3
    header = [c for c in t.cells if c.row == 0]
    assert [c.text for c in header] == ["Год", "Месяц", "Значения коэффициентов"]


def test_run_tesseract_text_marks_low_confidence_and_skips_table_area():
    # Взяты реальные порядки величин confidence из эксперимента со штампом
    # (докс.stage3-plan.md): чистый текст 71-97, повреждённый печатью 0-55.
    data = _tesseract_data(
        [
            ("Председатель", 96, 10, 10, 100, 20),
            ("Правительства", 1, 120, 10, 100, 20),  # под печатью
            ("Д.Медведев", 95, 230, 10, 100, 20),
            ("внутри_таблицы", 90, 10, 100, 50, 20),  # должно быть пропущено
        ]
    )
    table_bboxes = [(0.0, 90.0, 200.0, 130.0)]
    blocks = run_tesseract_text(
        Path("unused.png"), table_bboxes, image_data_fn=lambda _p: data, confidence_threshold=60
    )
    assert len(blocks) == 1
    assert blocks[0].text == "Председатель Правительства Д.Медведев"
    assert blocks[0].low_words == (False, True, False)


def test_expected_signatory_title_known_authorities():
    assert (
        expected_signatory_title("Правительство Российской Федерации")
        == "Председатель Правительства Российской Федерации"
    )
    assert (
        expected_signatory_title("Президент Российской Федерации")
        == "Президент Российской Федерации"
    )
    assert expected_signatory_title("Федеральная служба по аккредитации") == "Руководитель"
    # Министерства пока не поддержаны — склонение названия ведомства не выводится
    # надёжно без риска угадывания (см. docs/stage3-plan.md).
    assert expected_signatory_title("Министерство природных ресурсов") is None
    assert expected_signatory_title("Неизвестный орган") is None


def test_run_tesseract_text_restores_damaged_title_formula_without_touching_name():
    # Реальный кейс из эксперимента со штампом: формула должности и ФИО
    # подписанта стоят на одной строке ("Председатель Правительства Российской
    # Федерации Д.Медведев"), "Правительства" повреждено печатью — должна
    # восстановиться только формула, ФИО не трогается.
    data = _tesseract_data(
        [
            ("Председатель", 96, 10, 10, 100, 20),
            ("Правительства", 1, 120, 10, 100, 20),  # под печатью
            ("Российской", 90, 230, 10, 100, 20),
            ("Федерации", 93, 340, 10, 100, 20),
            ("Д.Медведев", 95, 450, 10, 100, 20),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Правительство Российской Федерации",
    )
    assert len(blocks) == 1
    assert blocks[0].text == "Председатель Правительства Российской Федерации Д.Медведев"


def test_run_tesseract_text_restores_title_split_across_tesseract_lines_near_stamp():
    # Реальный случай (живой тест, 12-стр. документ): печать сдвинула геометрию
    # слов так, что Tesseract разбил формулу должности на ДВЕ отдельные строки —
    # "Председатель [нрзб.]" / "Российской [нрзб.] Д.Медведев". Восстановление
    # должно сработать сквозным поиском по словам страницы, не только внутри
    # одной уже сгруппированной строки.
    data = _tesseract_data_multiline(
        [
            ("Председатель", 96, 10, 100, 100, 20, 1),
            ("Правительства", 1, 120, 100, 100, 20, 1),  # под печатью, строка 1
            ("Российской", 90, 10, 130, 100, 20, 2),
            ("Федерации", 1, 120, 130, 100, 20, 2),  # под печатью, строка 2
            ("Д.Медведев", 95, 230, 130, 100, 20, 2),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Правительство Российской Федерации",
    )
    assert len(blocks) == 2
    assert blocks[0].text == "Председатель Правительства"
    assert blocks[1].text == "Российской Федерации Д.Медведев"


def test_run_tesseract_text_fixes_high_confidence_misread_of_title_word():
    # Реальный случай (калибровка на 4 документах, 2026-10-02): печать на
    # «Федеральный закон №310-ФЗ» исказила «Президент» до «Ярезидент», но
    # Tesseract дал этому слову confidence=73 — выше порога 60, поэтому оно
    # НЕ было замаскировано как [нрзб.]. Восстановление должно сработать и на
    # уже "чистом" (по confidence), но текстово похожем на ожидаемое слове.
    data = _tesseract_data(
        [
            ("Ярезидент", 73, 10, 10, 150, 20),
            ("Российской", 95, 170, 10, 150, 20),
            ("Федерации", 96, 330, 10, 150, 20),
            ("В.Путин", 62, 490, 10, 150, 20),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Президент Российской Федерации",
    )
    assert len(blocks) == 1
    assert blocks[0].text == "Президент Российской Федерации В.Путин"


def test_run_tesseract_text_fixes_truncated_title_words_across_lines():
    # Реальный случай (Постановление №1430, 2026-10-02): печать обрезала слова
    # формулы ("Правительства"→"Прави", "Федерации"→"Федей"), конфиденс не
    # замаскировал их, полное нечёткое сравнение тоже не ловило (разная длина
    # режет ratio до 0.56/0.57) — добавлено отдельное префиксное сравнение.
    # Формула к тому же снова разорвана на две строки, как и в более раннем
    # реальном случае ("Председатель Прави" / "Российской Федей М.Мишустин").
    data = _tesseract_data_multiline(
        [
            ("Председатель", 95, 10, 100, 150, 20, 1),
            ("Прави", 65, 170, 100, 80, 20, 1),  # искажено печатью, но выше порога 60
            ("Российской", 90, 10, 130, 150, 20, 2),
            ("Федей", 68, 170, 130, 80, 20, 2),  # искажено печатью, но выше порога 60
            ("М.Мишустин", 93, 260, 130, 150, 20, 2),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Правительство Российской Федерации",
    )
    assert len(blocks) == 2
    assert blocks[0].text == "Председатель Правительства"
    assert blocks[1].text == "Российской Федерации М.Мишустин"


def test_run_tesseract_text_does_not_corrupt_unrelated_mentions_of_russian_federation():
    # "Российской Федерации" — обычная фраза, многократно встречающаяся в теле
    # любого НПА вне подписи. Слово перед ней не похоже на "Президент"/
    # "Председатель" — окно должно быть отклонено целиком, без порчи текста.
    data = _tesseract_data(
        [
            ("в", 95, 10, 10, 30, 20),
            ("соответствии", 95, 50, 10, 150, 20),
            ("с", 95, 210, 10, 30, 20),
            ("законодательством", 95, 250, 10, 220, 20),
            ("Российской", 95, 480, 10, 150, 20),
            ("Федерации", 95, 640, 10, 150, 20),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Президент Российской Федерации",
    )
    assert len(blocks) == 1
    assert blocks[0].text == "в соответствии с законодательством Российской Федерации"


def test_run_tesseract_text_does_not_touch_clean_lines_or_unrelated_damage():
    data = _tesseract_data(
        [
            ("Обычный", 95, 10, 10, 100, 20),
            ("абзац", 1, 120, 10, 100, 20),  # повреждено, но не формула должности
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"),
        [],
        image_data_fn=lambda _p: data,
        confidence_threshold=60,
        authority_name="Правительство Российской Федерации",
    )
    assert blocks[0].text == "Обычный абзац"
    assert blocks[0].low_words == (False, True)


def test_run_tesseract_text_without_authority_name_leaves_mark_as_is():
    data = _tesseract_data(
        [
            ("Председатель", 96, 10, 10, 100, 20),
            ("Правительства", 1, 120, 10, 100, 20),
        ]
    )
    blocks = run_tesseract_text(
        Path("unused.png"), [], image_data_fn=lambda _p: data, confidence_threshold=60
    )
    assert blocks[0].text == "Председатель Правительства"
    assert blocks[0].low_words == (False, True)


def test_ocr_page_orders_tables_and_text_by_vertical_position():
    fixture = _load_simple_table_json()
    text_data = _tesseract_data([("Заголовок", 95, 10, 5, 100, 20)])

    blocks = ocr_page(
        Path("unused.png"),
        predict_tables_fn=lambda _p: fixture,
        image_data_fn=lambda _p: text_data,
    )
    assert len(blocks) == 2
    assert isinstance(blocks[0], TextBlock)  # y=5, выше таблицы (y~1412 в фикстуре)
    assert isinstance(blocks[1], TableBlock)


def test_build_docx_writes_text_and_merged_table(tmp_path: Path):
    from docx import Document

    cells = [
        TableCell(row=0, col=0, row_span=1, col_span=2, text="Заголовок на всю ширину"),
        TableCell(row=1, col=0, row_span=1, col_span=1, text="a"),
        TableCell(row=1, col=1, row_span=1, col_span=1, text="b"),
    ]
    pages = [
        PageResult(page_number=1, ocr_applied=False, plain_text="Обычный текстовый слой"),
        PageResult(
            page_number=2,
            ocr_applied=True,
            blocks=[
                TextBlock(y=0, text="Распознанный абзац"),
                TableBlock(y=10, bbox=(0, 10, 100, 100), cells=cells, n_rows=2, n_cols=2),
            ],
        ),
    ]
    out = tmp_path / "out.docx"
    build_docx(pages, out)
    assert out.is_file()

    doc = Document(str(out))
    texts = [p.text for p in doc.paragraphs]
    assert "Обычный текстовый слой" in texts
    assert "Распознанный абзац" in texts
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert table.cell(0, 0).text == "Заголовок на всю ширину"
    # Объединение по горизонтали — (0,0) и (0,1) должны указывать на одну физическую ячейку.
    assert table.cell(0, 0)._tc is table.cell(0, 1)._tc
    assert table.cell(1, 0).text == "a"
    assert table.cell(1, 1).text == "b"


def test_build_docx_survives_non_rectangular_merge_conflict(tmp_path: Path, monkeypatch):
    # Реальный случай (Постановление №1430, 2026-10-02): на сложной реальной
    # таблице геометрия двух ячеек даёт конфликтующие span'ы, которые python-docx
    # не может представить одним прямоугольным объединением
    # (InvalidSpanError: requested span not rectangular). Раньше это роняло всю
    # сборку DOCX, теряя OCR всех страниц документа. Теперь такая ячейка остаётся
    # неотъединённой (текст не теряется), а не валит всю сборку.
    import docx.table
    from docx.exceptions import InvalidSpanError

    original_merge = docx.table._Cell.merge
    call_count = {"n": 0}

    def _flaky_merge(self, other_cell):
        call_count["n"] += 1
        if call_count["n"] == 1:
            raise InvalidSpanError("requested span not rectangular")
        return original_merge(self, other_cell)

    monkeypatch.setattr(docx.table._Cell, "merge", _flaky_merge)

    cells = [
        TableCell(row=0, col=0, row_span=1, col_span=2, text="Конфликтующая ячейка"),
        TableCell(row=1, col=0, row_span=1, col_span=1, text="a"),
        TableCell(row=1, col=1, row_span=1, col_span=1, text="b"),
    ]
    pages = [
        PageResult(
            page_number=1,
            ocr_applied=True,
            blocks=[TableBlock(y=0, bbox=(0, 0, 100, 100), cells=cells, n_rows=2, n_cols=2)],
        ),
    ]
    out = tmp_path / "out.docx"
    build_docx(pages, out)  # не должно бросить InvalidSpanError
    assert out.is_file()

    from docx import Document

    doc = Document(str(out))
    table = doc.tables[0]
    # Текст не потерян, даже хотя объединение не удалось.
    assert table.cell(0, 0).text == "Конфликтующая ячейка"
    assert table.cell(1, 0).text == "a"
    assert table.cell(1, 1).text == "b"


def test_ocr_document_reuses_existing_text_layer_without_calling_engines(tmp_path: Path):
    import fitz

    pdf_path = tmp_path / "with_text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    # Латиница намеренно — встроенный шрифт fitz.insert_text() по умолчанию не
    # поддерживает кириллицу (буквы превращаются в "." при извлечении текста);
    # для этого теста важна только сама логика пропуска OCR, не язык содержимого.
    page.insert_text((72, 72), "This page already has a real text layer, not a scan.")
    doc.save(str(pdf_path))
    doc.close()

    def _should_not_be_called(_path):
        raise AssertionError("движок не должен вызываться — у страницы есть текстовый слой")

    result = ocr_document(
        pdf_path,
        tmp_path / "out.docx",
        predict_tables_fn=_should_not_be_called,
        image_data_fn=_should_not_be_called,
    )
    assert isinstance(result, OcrDocumentResult)
    assert len(result.pages) == 1
    assert result.pages[0].ocr_applied is False
    assert "real text layer" in result.pages[0].plain_text
    assert result.docx_path.is_file()


def test_ocr_document_does_not_build_default_engines_when_no_scan_pages(
    tmp_path: Path, monkeypatch
):
    # PPStructureV3 грузит ~19 моделей (минуты, сотни МБ) — документ без единой
    # страницы-скана не должен платить эту цену; CLI smoke-тест на реальном
    # венве подтвердил это (0.3с вместо нескольких минут), здесь — то же самое
    # без реальных тяжёлых движков, через monkeypatch default-конструктора.
    import fitz

    import npa_pipeline.ocr as ocr_module

    def _should_not_be_called():
        raise AssertionError(
            "_default_predict_tables_fn не должен вызываться — у документа нет страниц-сканов"
        )

    monkeypatch.setattr(ocr_module, "_default_predict_tables_fn", _should_not_be_called)

    pdf_path = tmp_path / "with_text.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), "This page already has a real text layer, not a scan.")
    doc.save(str(pdf_path))
    doc.close()

    result = ocr_document(pdf_path, tmp_path / "out.docx")
    assert result.pages[0].ocr_applied is False


def test_ocr_document_runs_hybrid_pipeline_when_no_text_layer(tmp_path: Path):
    import fitz

    pdf_path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()  # пустая страница — имитация скана без текстового слоя
    doc.save(str(pdf_path))
    doc.close()

    fixture = _load_simple_table_json()
    text_data = _tesseract_data([("Распознано", 90, 10, 5, 100, 20)])
    calls = {"tables": 0, "text": 0}

    def fake_predict(_path):
        calls["tables"] += 1
        return fixture

    def fake_image_data(_path):
        calls["text"] += 1
        return text_data

    result = ocr_document(
        pdf_path,
        tmp_path / "out.docx",
        predict_tables_fn=fake_predict,
        image_data_fn=fake_image_data,
    )
    assert calls["tables"] == 1
    assert calls["text"] == 1
    assert result.pages[0].ocr_applied is True
    assert any(isinstance(b, TableBlock) for b in result.pages[0].blocks)
    assert result.docx_path.is_file()


def test_ensure_free_memory_raises_below_threshold(monkeypatch):
    import psutil

    from npa_pipeline.ocr import InsufficientMemoryError, ensure_free_memory

    class _Mem:
        available = 500 * 1024 * 1024

    monkeypatch.setattr(psutil, "virtual_memory", lambda: _Mem())
    with pytest.raises(InsufficientMemoryError, match="перед страницей 3"):
        ensure_free_memory(2048, "перед страницей 3")
    ensure_free_memory(0, "отключено")  # 0 — проверка отключена, ничего не делает


def test_ocr_document_stops_before_engines_when_memory_is_low(tmp_path: Path, monkeypatch):
    import fitz
    import psutil

    from npa_pipeline.ocr import InsufficientMemoryError

    class _Mem:
        available = 100 * 1024 * 1024

    monkeypatch.setattr(psutil, "virtual_memory", lambda: _Mem())
    pdf_path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    def _should_not_load(*_a, **_k):
        raise AssertionError("движок не должен загружаться при нехватке памяти")

    import npa_pipeline.ocr as ocr_module

    monkeypatch.setattr(ocr_module, "_default_predict_tables_fn", _should_not_load)
    with pytest.raises(InsufficientMemoryError):
        ocr_document(pdf_path, tmp_path / "out.docx", min_free_memory_mb=2048)


def test_text_paragraphs_restores_centered_lines_indents_and_breaks():
    from npa_pipeline.ocr import _text_paragraphs

    # Геометрия реальной первой страницы Постановления №1430 (200 DPI).
    run = [
        TextBlock(y=435, text="ПРАВИТЕЛЬСТВО РОССИЙСКОЙ ФЕДЕРАЦИИ", x0=270, x1=1381, height=53),
        TextBlock(y=821, text="Об утверждении технологических показателей", x0=313, x1=1337, height=35),
        TextBlock(y=866, text="доступных технологий в сфере", x0=366, x1=1281, height=34),
        TextBlock(y=1088, text="В соответствии с пунктом 5 статьи 23", x0=300, x1=1447, height=44),
        TextBlock(y=1142, text="Федерального закона Правительство", x0=203, x1=1446, height=35),
        TextBlock(y=1201, text="постановляет:", x0=202, x1=563, height=19),
        TextBlock(y=1242, text="Утвердить прилагаемые показатели", x0=299, x1=1447, height=34),
        TextBlock(y=1290, text="доступных технологий в сфере", x0=201, x1=1446, height=36),
    ]
    paras = _text_paragraphs(run, dpi=200)
    # «постановляет:» — короткая последняя строка вводной части, абзац не начинает.
    assert [p[0] for p in paras] == [
        "ПРАВИТЕЛЬСТВО РОССИЙСКОЙ ФЕДЕРАЦИИ",
        "Об утверждении технологических показателей",
        "доступных технологий в сфере",
        "В соответствии с пунктом 5 статьи 23 Федерального закона Правительство постановляет:",
        "Утвердить прилагаемые показатели доступных технологий в сфере",
    ]
    assert [p[1] for p in paras][:3] == ["center", "center", "center"]
    assert paras[3][1] == "left" and paras[3][2] > 0  # «В соответствии» начинается с отступа (x0=300)
    assert paras[4][2] > 0  # «Утвердить» с отступом первой строки


def test_text_paragraphs_without_geometry_keeps_one_line_per_paragraph():
    from npa_pipeline.ocr import _text_paragraphs

    run = [TextBlock(y=0, text="a"), TextBlock(y=40, text="b")]
    assert [(t, a, f) for t, a, f, _ in _text_paragraphs(run, dpi=200)] == [("a", "left", 0.0), ("b", "left", 0.0)]


def test_remove_running_headers_and_footers_drops_page_numbers_and_repeats_only():
    from npa_pipeline.ocr import remove_running_headers_and_footers

    def page(n: int, footer_text: str, body: list[str]) -> PageResult:
        blocks = [TextBlock(y=500 + 60 * i, text=t, x0=200, x1=1400, height=35) for i, t in enumerate(body)]
        blocks.append(TextBlock(y=960, text=footer_text, x0=200, x1=400, height=20))  # нижний колонтитул
        blocks.append(TextBlock(y=20, text=str(n), x0=800, x1=850, height=20))  # номер страницы сверху
        return PageResult(page_number=n, ocr_applied=True, blocks=blocks, page_height=1000.0)

    pages = [
        page(1, "4678924", ["Текст первой страницы"]),
        page(2, "4678924", ["Текст второй страницы"]),
        page(3, "4678924", ["Текст третьей страницы"]),
    ]
    cleaned = remove_running_headers_and_footers(pages)
    for p in cleaned:
        texts = [b.text for b in p.blocks]
        assert "4678924" not in texts
        assert str(p.page_number) not in texts
    assert [b.text for b in cleaned[0].blocks] == ["Текст первой страницы"]


def test_remove_running_headers_keeps_repeated_text_in_body_and_single_footer_lines():
    from npa_pipeline.ocr import remove_running_headers_and_footers

    body_repeat = TextBlock(y=500, text="Одинаковая строка тела", x0=200, x1=1400, height=35)
    pages = [
        PageResult(page_number=i, ocr_applied=True, blocks=[body_repeat, TextBlock(y=960, text=f"Колонтитул {name}", x0=200, x1=600, height=20)], page_height=1000.0)
        for i, name in ((1, "alpha"), (2, "beta"), (3, "gamma"))
    ]
    cleaned = remove_running_headers_and_footers(pages)
    for p in cleaned:
        texts = [b.text for b in p.blocks]
        assert "Одинаковая строка тела" in texts  # вне полос — не трогаем, даже если повторяется
        assert any(t.startswith("Колонтитул ") for t in texts)  # уникален — не повторяется, сохраняем


def test_run_tatr_tables_builds_block_with_offset_and_text_from_ocr(tmp_path: Path):
    from PIL import Image

    from npa_pipeline.ocr import run_tatr_tables
    from npa_pipeline.table_structure import TableStructure

    img_path = tmp_path / "page.png"
    Image.new("RGB", (600, 600), "white").save(img_path)

    def fake_region(_img):
        return [(100.0, 100.0, 400.0, 300.0)]

    def fake_structure(_crop):
        # сетка 2×2 в координатах кропа; третьего столбца нет, он пуст
        return TableStructure(
            rows=[(0.0, 0.0, 300.0, 100.0), (0.0, 100.0, 300.0, 200.0)],
            cols=[(0.0, 0.0, 150.0, 200.0), (150.0, 0.0, 300.0, 200.0), (300.0, 0.0, 320.0, 200.0)],
            spans=[],
        )

    def fake_text(_crop_path):
        return {
            "table_res_list": [
                {
                    "table_ocr_pred": {
                        "rec_boxes": [[10, 10, 140, 40], [160, 110, 290, 140]],
                        "rec_texts": ["Ячейка-1", "Ячейка-2"],
                    }
                }
            ]
        }

    blocks = run_tatr_tables(
        img_path, fake_text, region_fn=fake_region, structure_fn=fake_structure
    )
    assert len(blocks) == 1
    b = blocks[0]
    assert b.dropped_columns == 1  # пустой третий столбец удалён и учтён
    assert (b.n_rows, b.n_cols) == (2, 2)
    by_pos = {(c.row, c.col): c.text for c in b.cells}
    assert by_pos[(0, 0)] == "Ячейка-1"
    assert by_pos[(1, 1)] == "Ячейка-2"
    x1, y1, x2, y2 = b.bbox
    assert 80.0 <= x1 and x2 <= 420.0 and 80.0 <= y1 and y2 <= 320.0  # bbox — по содержимому, с отступом вокруг текста


def test_run_tatr_tables_drops_duplicate_texts_from_overlapping_table_entries(tmp_path: Path):
    from PIL import Image

    from npa_pipeline.ocr import run_tatr_tables
    from npa_pipeline.table_structure import TableStructure

    img_path = tmp_path / "page.png"
    Image.new("RGB", (600, 600), "white").save(img_path)

    def fake_text(_crop_path):
        # две записи таблиц на одном кропе содержат одну и ту же строку
        same = {"rec_boxes": [[10, 10, 140, 40]], "rec_texts": ["Ячейка"]}
        return {"table_res_list": [{"table_ocr_pred": same}, {"table_ocr_pred": same}]}

    blocks = run_tatr_tables(
        img_path,
        fake_text,
        region_fn=lambda _img: [(100.0, 100.0, 400.0, 300.0)],
        structure_fn=lambda _crop: TableStructure(
            rows=[(0.0, 0.0, 300.0, 100.0)], cols=[(0.0, 0.0, 150.0, 100.0), (150.0, 0.0, 300.0, 100.0)], spans=[]
        ),
    )
    assert blocks[0].cells[0].text == "Ячейка"  # не "Ячейка Ячейка"


def test_run_tatr_tables_recognizes_each_cell_separately_when_cell_fn_given(tmp_path: Path):
    from PIL import Image

    from npa_pipeline.ocr import run_tatr_tables
    from npa_pipeline.table_structure import TableStructure

    img_path = tmp_path / "page.png"
    Image.new("RGB", (600, 600), "white").save(img_path)
    seen: list[str] = []

    def fake_cell_text(cell_path):
        seen.append(Path(cell_path).name)
        return "ЯЧ", (False,)

    blocks = run_tatr_tables(
        img_path,
        lambda _p: {"table_res_list": [{"table_ocr_pred": {"rec_boxes": [[10, 10, 90, 40]], "rec_texts": ["x"]}}]},
        cell_text_fn=fake_cell_text,
        region_fn=lambda _img: [(100.0, 100.0, 400.0, 300.0)],
        structure_fn=lambda _crop: TableStructure(
            rows=[(0.0, 0.0, 300.0, 100.0)],
            cols=[(0.0, 0.0, 150.0, 100.0), (150.0, 0.0, 300.0, 100.0)],
            spans=[],
        ),
    )
    assert seen  # каждая ячейка с текстом распознана отдельным кропом
    assert blocks[0].cells[0].text == "ЯЧ"


def test_unique_text_blocks_merges_same_text_with_shifted_boxes():
    from npa_pipeline.ocr import _unique_text_blocks

    text_json = {
        "table_res_list": [
            {"table_ocr_pred": {"rec_boxes": [[10, 10, 100, 40]], "rec_texts": ["N п/п"]}},
            {"table_ocr_pred": {"rec_boxes": [[12, 13, 102, 44]], "rec_texts": ["N п/п"]}},
            {"table_ocr_pred": {"rec_boxes": [[300, 10, 400, 40]], "rec_texts": ["N п/п"]}},
        ]
    }
    blocks = _unique_text_blocks(text_json)
    assert len(blocks) == 2  # сдвинутый дубликат убран; тот же текст в другом месте — оставлен


def test_run_tatr_tables_downscales_large_crop_for_text_and_restores_coordinates(tmp_path: Path):
    from PIL import Image

    from npa_pipeline.ocr import TEXT_PREDICT_MAX_SIDE, run_tatr_tables
    from npa_pipeline.table_structure import TableStructure

    img_path = tmp_path / "page.png"
    Image.new("RGB", (4000, 3500), "white").save(img_path)
    seen: dict[str, tuple[int, int]] = {}

    def fake_text(crop_path):
        seen["size"] = Image.open(crop_path).size
        return {"table_res_list": [{"table_ocr_pred": {"rec_boxes": [[100, 100, 300, 140]], "rec_texts": ["ок"]}}]}

    run_tatr_tables(
        img_path,
        fake_text,
        region_fn=lambda _img: [(0.0, 0.0, 3900.0, 3300.0)],
        structure_fn=lambda _crop: TableStructure(
            rows=[(0.0, 0.0, 4000.0, 4000.0)], cols=[(0.0, 0.0, 4000.0, 4000.0)], spans=[]
        ),
    )
    assert max(seen["size"]) <= TEXT_PREDICT_MAX_SIDE


def test_run_tatr_tables_flags_low_confidence_cell_and_docx_highlights_it(tmp_path: Path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX
    from PIL import Image

    from npa_pipeline.ocr import PageResult, TableBlock, build_docx, run_tatr_tables
    from npa_pipeline.table_structure import TableStructure

    img_path = tmp_path / "page.png"
    Image.new("RGB", (600, 600), "white").save(img_path)

    def fake_cell_text(cell_path):
        return ("сомнительно", (True,)) if "cell_0_0" in Path(cell_path).name else ("верно", (False,))

    blocks = run_tatr_tables(
        img_path,
        lambda _p: {"table_res_list": [{"table_ocr_pred": {"rec_boxes": [[10, 10, 90, 40], [160, 10, 240, 40]], "rec_texts": ["a", "b"]}}]},
        cell_text_fn=fake_cell_text,
        region_fn=lambda _img: [(100.0, 100.0, 400.0, 300.0)],
        structure_fn=lambda _crop: TableStructure(
            rows=[(0.0, 0.0, 300.0, 100.0)],
            cols=[(0.0, 0.0, 150.0, 100.0), (150.0, 0.0, 300.0, 100.0)],
            spans=[],
        ),
    )
    by_pos = {(c.row, c.col): c for c in blocks[0].cells}
    assert by_pos[(0, 0)].low_lines == (True,)
    assert not any(by_pos[(0, 1)].low_lines)

    out = tmp_path / "out.docx"
    build_docx([PageResult(page_number=1, ocr_applied=True, blocks=blocks)], out)
    doc = Document(str(out))
    table = doc.tables[0]
    highlighted = [r.text for r in table.cell(0, 0).paragraphs[0].runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert highlighted == ["сомнительно"]
    assert all(r.font.highlight_color is None for r in table.cell(0, 1).paragraphs[0].runs)
    assert any("низкой уверенностью" in p.text for p in doc.paragraphs)


def test_body_text_low_confidence_word_is_highlighted_in_docx(tmp_path: Path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    from npa_pipeline.ocr import PageResult, build_docx

    data = _tesseract_data([("Текст", 95, 10, 10, 100, 20), ("сомнительно", 20, 120, 10, 100, 20)])
    blocks = run_tesseract_text(Path("unused.png"), [], image_data_fn=lambda _p: data, confidence_threshold=60)
    out = tmp_path / "body.docx"
    build_docx([PageResult(page_number=1, ocr_applied=True, blocks=blocks)], out)
    doc = Document(str(out))
    para = next(p for p in doc.paragraphs if "Текст" in p.text)
    highlighted = [r.text for r in para.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert highlighted == ["сомнительно"]
    assert "[нрзб.]" not in para.text


def _two_text_pages_pdf(path: Path) -> Path:
    import fitz

    doc = fitz.open()
    for _ in range(2):
        doc.new_page().insert_text((72, 72), "This page already has a real text layer, not a scan.")
    doc.save(str(path))
    doc.close()
    return path


def test_ocr_document_reports_progress_and_saves_each_page(tmp_path: Path):
    pdf_path = _two_text_pages_pdf(tmp_path / "two.pdf")
    events = []
    ocr_document(
        pdf_path,
        tmp_path / "out.docx",
        progress_fn=lambda done, total, n, status: events.append((done, total, n, status)),
    )
    assert events == [(1, 2, 1, "текстовый слой"), (2, 2, 2, "текстовый слой")]
    pages_dir = tmp_path / "out.pages"
    assert sorted(p.name for p in pages_dir.iterdir()) == ["page_0001.json", "page_0002.json"]


def test_ocr_document_resumes_from_saved_pages_without_engines(tmp_path: Path):
    pdf_path = _two_text_pages_pdf(tmp_path / "two.pdf")
    ocr_document(pdf_path, tmp_path / "out.docx", progress_fn=lambda *a: None)

    def _should_not_be_called(_path):
        raise AssertionError("страницы уже сохранены — движки не должны вызываться")

    events = []
    result = ocr_document(
        pdf_path,
        tmp_path / "out.docx",
        predict_tables_fn=_should_not_be_called,
        image_data_fn=_should_not_be_called,
        progress_fn=lambda done, total, n, status: events.append(status),
    )
    assert [p.page_number for p in result.pages] == [1, 2]
    assert events == ["взято из сохранённого прогона"] * 2


def test_ocr_document_ignores_saved_pages_with_other_parameters(tmp_path: Path):
    pdf_path = _two_text_pages_pdf(tmp_path / "two.pdf")
    ocr_document(pdf_path, tmp_path / "out.docx", dpi=200, progress_fn=lambda *a: None)
    events = []
    ocr_document(
        pdf_path,
        tmp_path / "out.docx",
        dpi=150,
        progress_fn=lambda done, total, n, status: events.append(status),
    )
    assert events == ["текстовый слой"] * 2


def test_saved_page_roundtrip_keeps_text_and_table_blocks(tmp_path: Path):
    from npa_pipeline.ocr import _load_page, _save_page

    table = TableBlock(
        y=10.0,
        bbox=(1.0, 10.0, 200.0, 90.0),
        cells=[TableCell(row=0, col=0, row_span=1, col_span=1, text="а\nб", low_lines=(False, True))],
        n_rows=1,
        n_cols=1,
        dropped_columns=2,
    )
    text = TextBlock(y=100.0, text="Обычный текст", x0=5.0, x1=300.0, height=20.0, low_words=(False, True))
    page = PageResult(page_number=3, ocr_applied=True, blocks=[table, text], page_height=1200.0)
    meta = {"dpi": 200}
    path = tmp_path / "page_0003.json"
    _save_page(path, page, meta)

    loaded = _load_page(path, meta)
    assert loaded == page
    assert _load_page(path, {"dpi": 150}) is None


def test_ocr_document_writes_timings_per_document_and_per_page(tmp_path: Path):
    pdf_path = _two_text_pages_pdf(tmp_path / "two.pdf")
    result = ocr_document(pdf_path, tmp_path / "out.docx", progress_fn=lambda *a: None)
    timings = json.loads(result.timings_path.read_text(encoding="utf-8"))
    assert result.timings_path == tmp_path / "out.timings.json"
    assert [s["step"] for s in timings["document_steps"]] == [
        "Открытие и подготовка",
        "Завершение распознавания",
        "Колонтитулы",
        "Сборка DOCX",
    ]
    assert sorted(timings["pages"]) == ["1", "2"]
    for page in timings["pages"].values():
        assert [s["step"] for s in page["steps"]][0] == "проверка сохранённого результата"
        assert page["total_seconds"] >= 0
    assert timings["total_seconds"] >= sum(s["seconds"] for s in timings["document_steps"])


def test_lazy_engine_loads_once_and_reloads_after_release():
    from npa_pipeline.ocr import _LazyEngine

    created = []

    def factory():
        created.append(1)
        return lambda x: x * 2

    engine = _LazyEngine(factory)
    assert engine(3) == 6
    assert engine(4) == 8
    assert len(created) == 1
    engine.release()
    assert engine(1) == 2
    assert len(created) == 2
