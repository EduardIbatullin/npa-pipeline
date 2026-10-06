"""Живой прогон ocr_document на реальном скане — настоящие PaddleOCR + Tesseract,
не фейки (см. tests/unit/test_ocr.py для изолированных тестов на фикстурах).
Медленно (минуты на страницу) — не входит в обычный прогон pytest."""

from __future__ import annotations

import re
from pathlib import Path

import pytest

from npa_pipeline.ocr import ocr_document

REAL_SCAN = (
    Path(__file__).parents[2]
    / "downloads"
    / "Постановление Правительства Российской Федерации от 31_05_2019 N 691.pdf"
)


@pytest.mark.live
def test_ocr_document_real_stamped_document(tmp_path: Path):
    if not REAL_SCAN.is_file():
        pytest.skip(f"нет реального файла для теста: {REAL_SCAN}")

    out_docx = tmp_path / "out.docx"
    result = ocr_document(
        REAL_SCAN, out_docx, authority_name="Правительство Российской Федерации"
    )

    assert result.docx_path.is_file()
    assert len(result.pages) == 12
    assert all(p.ocr_applied for p in result.pages)  # документ полностью скан

    from docx import Document

    doc = Document(str(out_docx))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "ПОСТАНОВЛЕНИЕ" in full_text
    assert "Правительства Российской Федерации" in full_text
    # Формула должности подписанта должна восстановиться, несмотря на печать (см.
    # docs/stage3-plan.md, «Реализация»). Пробел между словами допускается любой,
    # включая перевод строки — реальный документ разрывает эту формулу на две
    # отдельные строки Tesseract из-за геометрической помехи от штампа;
    # восстановление работает по словам страницы, не по одной строке (см.
    # _restore_signatory_title_words), поэтому слова корректны, но граница
    # строки/абзаца между ними сохраняется как у исходного скана.
    assert re.search(
        r"Председатель\s+Правительства\s+Российской\s+Федерации", full_text
    )


REAL_SCAN_PRESIDENT = (
    Path(__file__).parents[2] / "downloads" / "Федеральный закон от 31_07_2025 N 310-ФЗ.pdf"
)


@pytest.mark.live
def test_ocr_document_fixes_high_confidence_misread_of_president_title(tmp_path: Path):
    # Реальный случай из калибровки (docs/stage3-plan.md, «Калибровка порога»):
    # печать на последней странице этого документа искажает «Президент» до
    # «Ярезидент» с confidence=73 — выше порога, т.е. НЕ маскируется как [нрзб.],
    # и раньше проходило бы в вывод неисправленным. Весь 40-страничный документ
    # гонять ради одной строки слишком долго — вырезаем только последнюю
    # страницу в отдельный PDF, прогон остаётся настоящим (реальный скан,
    # реальные движки), просто короче.
    if not REAL_SCAN_PRESIDENT.is_file():
        pytest.skip(f"нет реального файла для теста: {REAL_SCAN_PRESIDENT}")

    import pypdf

    reader = pypdf.PdfReader(str(REAL_SCAN_PRESIDENT))
    writer = pypdf.PdfWriter()
    writer.add_page(reader.pages[-1])
    last_page_pdf = tmp_path / "last_page.pdf"
    with open(last_page_pdf, "wb") as f:
        writer.write(f)

    out_docx = tmp_path / "out.docx"
    result = ocr_document(
        last_page_pdf, out_docx, authority_name="Президент Российской Федерации"
    )

    assert result.pages[0].ocr_applied is True

    from docx import Document

    doc = Document(str(out_docx))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "Ярезидент" not in full_text
    assert re.search(r"Президент\s+Российской\s+Федерации", full_text)
    assert "В.Путин" in full_text
